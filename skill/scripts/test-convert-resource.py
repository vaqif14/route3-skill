#!/usr/bin/env python3
"""Real local document-conversion checks in a temporary workspace."""

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile


SCRIPT = Path(__file__).with_name("convert-resource.py")
RUNTIME = Path.home() / ".local/share/route3/markitdown/bin/python"


def pdf_bytes(text=None):
    """A complete one-page PDF; no text yields a raster-only page."""
    image = text is None
    content = (b"q 100 0 0 100 50 50 cm /Im1 Do Q" if image else
               f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
    resources = b"/XObject << /Im1 6 0 R >>" if image else b"/Font << /F1 4 0 R >>"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << "
               + resources + b" >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream"]
    if image:
        objects.append(b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
                       b"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\nstream\n"
                       b"\xff\xff\xff\nendstream")
    body = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(body))
        body += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    body += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return body


class ResourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="route3-resource-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def invoke(self, source, interpreter=None):
        proc = subprocess.run([str(interpreter or RUNTIME), str(SCRIPT), str(source),
                               "--cwd", str(self.root)], capture_output=True, text=True,
                              timeout=40)
        if proc.returncode in (0, 1):
            return proc.returncode, json.loads(proc.stdout)
        self.assertEqual(proc.returncode, 2, proc.stderr)
        return proc.returncode, json.loads(proc.stderr.strip().splitlines()[-1])

    def sample(self, suffix=".pdf"):
        source = self.root / f"resource with spaces{suffix}"
        source.write_bytes(pdf_bytes("Route3 sample value 42."))
        return source

    def test_text_pdf_converts_without_changing_source(self):
        source = self.sample()
        before = source.read_bytes()
        code, report = self.invoke(source)
        self.assertEqual(code, 0, report)
        self.assertIn("value 42", Path(report["markdown_path"]).read_text())
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(report["source_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(report["identity"]["converter_version"], "0.1.8")

    def test_cache_reused_and_corrupted_markdown_rebuilt(self):
        source = self.sample()
        _, first = self.invoke(source)
        _, second = self.invoke(source)
        self.assertFalse(first["cache_hit"])
        self.assertTrue(second["cache_hit"])
        Path(first["markdown_path"]).write_text("corrupted")
        _, third = self.invoke(source)
        self.assertFalse(third["cache_hit"])
        self.assertIn("value 42", Path(third["markdown_path"]).read_text())

    def test_changed_source_gets_new_artifact(self):
        source = self.sample()
        _, first = self.invoke(source)
        source.write_bytes(pdf_bytes("Route3 new value 84."))
        code, second = self.invoke(source)
        self.assertEqual(code, 0)
        self.assertNotEqual(first["markdown_path"], second["markdown_path"])
        self.assertIn("value 84", Path(second["markdown_path"]).read_text())
        self.assertTrue(Path(first["markdown_path"]).exists())

    def test_raster_pdf_requires_ocr(self):
        source = self.root / "raster-only.pdf"
        source.write_bytes(pdf_bytes())
        code, report = self.invoke(source)
        self.assertEqual(code, 1, report)
        self.assertEqual(report["status"], "needs_ocr")
        self.assertFalse(Path(report["markdown_path"]).read_text().strip())

    def test_uppercase_pdf_and_blank_text_status(self):
        source = self.sample(".PDF")
        code, report = self.invoke(source)
        self.assertEqual(code, 0, report)
        blank = self.root / "whitespace.txt"
        blank.write_text("  \n\t ")
        code, report = self.invoke(blank)
        self.assertEqual(code, 1, report)
        self.assertEqual(report["status"], "needs_review")

    def test_docx_real_conversion(self):
        source = self.root / "notes.docx"
        with ZipFile(source, "w") as archive:
            archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Route3 document value 42</w:t></w:r></w:p></w:body></w:document>')
        code, report = self.invoke(source)
        self.assertEqual(code, 0, report)
        self.assertIn("Route3 document value 42", Path(report["markdown_path"]).read_text())

    def test_pptx_real_conversion(self):
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Route3 slide"
        slide.placeholders[1].text = "Slide value 42"
        source = self.root / "deck.pptx"
        deck.save(source)
        code, report = self.invoke(source)
        self.assertEqual(code, 0, report)
        body = Path(report["markdown_path"]).read_text()
        self.assertIn("Route3 slide", body)
        self.assertIn("Slide value 42", body)

    def test_xlsx_real_conversion(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Route3Sheet"
        book.active.append(["Item", "Value"])
        book.active.append(["Sample", 42])
        source = self.root / "sheet.xlsx"
        book.save(source)
        code, report = self.invoke(source)
        self.assertEqual(code, 0, report)
        body = Path(report["markdown_path"]).read_text()
        self.assertIn("Route3Sheet", body)
        self.assertIn("42", body)
        self.assertTrue(any("formula" in note for note in report["warnings"]))

    def test_common_text_resources(self):
        for suffix, text in [(".html", "<h1>Route3</h1><p>Value 42</p>"),
                             (".csv", "Name,Value\nRoute3,42\n"),
                             (".json", '{"name":"Route3","value":42}'),
                             (".txt", "Route3 text value 42")]:
            with self.subTest(suffix=suffix):
                source = self.root / f"data{suffix}"
                source.write_text(text)
                code, report = self.invoke(source)
                self.assertEqual(code, 0, report)
                self.assertIn("Route3", Path(report["markdown_path"]).read_text())

    def test_existing_markdown_reused_without_artifacts(self):
        source = self.root / "notes.md"
        source.write_text("# Route3\n\nExisting Markdown.\n")
        code, report = self.invoke(source)
        self.assertEqual(code, 0)
        self.assertEqual(report["status"], "already_markdown")
        self.assertEqual(report["markdown_path"], str(source.resolve()))
        self.assertFalse((self.root / ".workflow").exists())

    def test_image_is_never_mistaken_for_full_text(self):
        from PIL import Image
        source = self.root / "image.png"
        Image.new("RGB", (10, 10), color="white").save(source)
        code, report = self.invoke(source)
        self.assertEqual(code, 1, report)
        self.assertEqual(report["status"], "needs_visual_review")

    def test_wrong_cached_paths_are_not_reused(self):
        source = self.sample()
        _, report = self.invoke(source)
        meta = Path(report["metadata_path"])
        body = json.loads(meta.read_text())
        body["markdown_path"] = "/incorrect/location.md"
        meta.write_text(json.dumps(body))
        code, regenerated = self.invoke(source)
        self.assertEqual(code, 0)
        self.assertFalse(regenerated["cache_hit"])
        self.assertTrue(Path(regenerated["markdown_path"]).exists())

    def test_corrupt_cache_status_cannot_promote_blank_content(self):
        source = self.root / "blank.txt"
        source.write_text(" \n ")
        code, report = self.invoke(source)
        self.assertEqual(code, 1)
        meta = Path(report["metadata_path"])
        body = json.loads(meta.read_text())
        body["status"] = "converted"
        meta.write_text(json.dumps(body))
        code, refreshed = self.invoke(source)
        self.assertEqual(code, 1, refreshed)
        self.assertEqual(refreshed["status"], "needs_review")
        self.assertFalse(refreshed["cache_hit"])

    def test_corrupt_cache_cannot_hide_extraction_warnings(self):
        source = self.sample()
        _, report = self.invoke(source)
        meta = Path(report["metadata_path"])
        body = json.loads(meta.read_text())
        body["warnings"] = []
        meta.write_text(json.dumps(body))
        code, refreshed = self.invoke(source)
        self.assertEqual(code, 0, refreshed)
        self.assertTrue(refreshed["warnings"])
        self.assertFalse(refreshed["cache_hit"])

    def test_corrupt_top_level_provenance_is_not_reused(self):
        source = self.sample()
        for field, value in [("source_path", "/wrong/source.pdf"),
                             ("source_sha256", "0" * 64), ("schema_version", 99)]:
            with self.subTest(field=field):
                _, report = self.invoke(source)
                meta = Path(report["metadata_path"])
                body = json.loads(meta.read_text())
                body[field] = value
                meta.write_text(json.dumps(body))
                code, refreshed = self.invoke(source)
                self.assertEqual(code, 0, refreshed)
                self.assertFalse(refreshed["cache_hit"])
                self.assertEqual(refreshed["source_path"], str(source.resolve()))
                self.assertEqual(refreshed["source_sha256"], hashlib.sha256(source.read_bytes()).hexdigest())

    def test_other_image_formats_are_explicitly_unsupported(self):
        from PIL import Image
        for suffix in [".gif", ".webp", ".bmp", ".tiff"]:
            with self.subTest(suffix=suffix):
                source = self.root / f"image{suffix}"
                Image.new("RGB", (10, 10), color="white").save(source)
                code, report = self.invoke(source)
                self.assertEqual(code, 2, report)
                self.assertIn("unsupported", report["message"])

    def test_failures_do_not_become_successful_markdown(self):
        for name, data in [("empty.pdf", b""), ("broken.pdf", b"not a PDF"),
                           ("archive.zip", b"archive is out of scope")]:
            with self.subTest(name=name):
                source = self.root / name
                source.write_bytes(data)
                code, report = self.invoke(source)
                self.assertEqual(code, 2, report)
                self.assertEqual(report["status"], "failed")
                self.assertEqual(source.read_bytes(), data)
        code, _ = self.invoke("https://example.invalid/resource.pdf")
        self.assertEqual(code, 2)
        code, _ = self.invoke(self.root)
        self.assertEqual(code, 2)

    def test_default_python_bootstraps_pinned_runtime(self):
        source = self.sample()
        code, report = self.invoke(source, interpreter=shutil.which("python3") or sys.executable)
        self.assertEqual(code, 0, report)
        self.assertEqual(report["identity"]["converter_version"], "0.1.8")

    def test_artifacts_have_private_permissions(self):
        _, report = self.invoke(self.sample())
        for key in ["markdown_path", "metadata_path"]:
            self.assertEqual(Path(report[key]).stat().st_mode & 0o777, 0o600)
        self.assertEqual(Path(report["markdown_path"]).parent.stat().st_mode & 0o777, 0o700)


if __name__ == "__main__":
    unittest.main()
